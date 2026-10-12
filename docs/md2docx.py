"""Convert the project's Markdown documents (SRS, SAD, Test Plan) to .docx.

Usage (from the repo root):  python docs/md2docx.py
Supports the Markdown subset used in these documents: headings, paragraphs,
bold/italic/inline code, bullet lists, tables, fenced code blocks and images.
"""
import os
import re
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
JOBS = [
    ("SRS.md", "docs/SRS.docx"),
    ("docs/SAD.md", "docs/SAD.docx"),
    ("docs/TEST_PLAN.md", "docs/Test_Plan.docx"),
]
INLINE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*\s][^*]*\*)")
LINK = re.compile(r"\[([^\]]+)\]\([^)]+\)")


def shade(cell, fill):
    pr = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd")
    el.set(qn("w:val"), "clear")
    el.set(qn("w:color"), "auto")
    el.set(qn("w:fill"), fill)
    pr.append(el)


def add_runs(par, text, size=None, bold=False):
    text = LINK.sub(r"\1", text)
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
        elif part.startswith("**") and part.endswith("**"):
            run = par.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            run = par.add_run(part[1:-1])
            run.italic = True
        else:
            run = par.add_run(part)
        if bold:
            run.bold = True
        if size:
            run.font.size = Pt(size)


def split_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Table Grid"
    for i, text in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.text = ""
        add_runs(cell.paragraphs[0], text, size=9, bold=True)
        shade(cell, "D9E2F3")
    for r in body:
        r = (r + [""] * ncols)[:ncols]
        cells = table.add_row().cells
        for i, text in enumerate(r):
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], text, size=9)
    doc.add_paragraph()


def convert(src_rel, dst_rel):
    src = os.path.join(ROOT, src_rel)
    base = os.path.dirname(src)
    lines = open(src, encoding="utf-8").read().splitlines()

    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    i = 0
    first_h1 = True
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if not stripped or stripped == "---" or stripped.startswith("<!--"):
            i += 1
            continue

        if stripped.startswith("```"):                      # code block
            i += 1
            code = []
            while i < len(lines) and not lines[i].strip().startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            for c in code:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.left_indent = Cm(0.5)
                run = p.add_run(c if c else " ")
                run.font.name = "Consolas"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
                run.font.size = Pt(8.5)
            doc.add_paragraph()
            continue

        m = re.match(r"^(#{1,4})\s+(.*)", stripped)         # heading
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1 and first_h1:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(LINK.sub(r"\1", text))
                run.bold = True
                run.font.size = Pt(20)
                first_h1 = False
            else:
                h = doc.add_heading(level=min(level, 4))
                add_runs(h, text)
            i += 1
            continue

        m = re.match(r"^!\[([^\]]*)\]\(([^)]+)\)", stripped)  # image
        if m:
            path = os.path.normpath(os.path.join(base, m.group(2)))
            if not os.path.exists(path):                      # path relative to repo root
                path = os.path.normpath(os.path.join(ROOT, m.group(2)))
            width = 4.6 if "component" in path else 6.3
            doc.add_picture(path, width=Inches(width))
            doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        if stripped.startswith("|"):                          # table
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_row(lines[i])
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells if c) or not any(cells):
                    if any(cells):
                        rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if re.match(r"^[-*]\s+", stripped):                   # bullet
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, re.sub(r"^[-*]\s+", "", stripped))
            i += 1
            continue

        indent = len(line) - len(line.lstrip(" "))            # plain paragraph
        p = doc.add_paragraph()
        if indent:
            p.paragraph_format.left_indent = Cm(0.5 * (indent // 3 or 1))
        add_runs(p, stripped)
        i += 1

    out = os.path.join(ROOT, dst_rel)
    doc.save(out)
    n_tables, n_pics = len(doc.tables), len(doc.inline_shapes)
    print(f"{dst_rel}: {len(doc.paragraphs)} paragraphs, {n_tables} tables, {n_pics} images")


if __name__ == "__main__":
    for s, d in JOBS:
        convert(s, d)
    sys.exit(0)
